"""Human judgments of Scripture Forge drafts from the 2025 field tests (the NT field test and a second trial).

The workbooks are exported from Google Drive with rclone (README). Teams filled in sheets of four kinds, each with a
header row naming its columns, in English, French or Indonesian:

    usability    a draft verse, and whether the team would correct it or start over from the source Bible
    comparison   two drafts of a verse from different models, and which is better (or both equally good or poor)
    correction   a draft verse and the team's corrected version
    human_ai     the team's own translation of a verse beside an AI draft, and whether the draft would have helped

and a fifth from the compiled feedback workbooks:

    analysis     a draft verse rated on the three criteria (same meaning, clear, sounds right), whether it is hard to
                 read aloud or confusing, which error types it has, and a score out of 100; the drafts come from the
                 teams' verse selection sheets, and the team later corrected the same verses

Each judged verse becomes one row, with the team's project and the verse reference, so references can be looked up.
"""

import argparse
import hashlib
import logging
import re
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional, Tuple

import pandas as pd

from ..scripture.verses import expand_vref

LOGGER = logging.getLogger(__name__)

BOOKS = {
    "GEN": "GEN",
    "GENESIS": "GEN",
    "JON": "JON",
    "RUT": "RUT",
    "MAT": "MAT",
    "MATTHEW": "MAT",
    "MATIUS": "MAT",
    "MATTHIEU": "MAT",
    "MRK": "MRK",
    "MARK": "MRK",
    "MARC": "MRK",
    "MARKUS": "MRK",
    "LUK": "LUK",
    "LUKE": "LUK",
    "LUC": "LUK",
    "LUKAS": "LUK",
    "JHN": "JHN",
    "JOHN": "JHN",
    "JEAN": "JHN",
    "YOHANES": "JHN",
    "ACT": "ACT",
    "ACTS": "ACT",
    "ACTES": "ACT",
    "KISAH": "ACT",
    "ROM": "ROM",
    "ROMANS": "ROM",
    "1CO": "1CO",
    "2CO": "2CO",
    "HEB": "HEB",
    "COL": "COL",
    "1TH": "1TH",
    "1TI": "1TI",
    "2TI": "2TI",
}

_DRAFT = re.compile(r"^(draft|draf|ebauche|ébauche|verset traduit|ai draft)\b", re.IGNORECASE)
_DRAFT_A = re.compile(r"^(draft|draf|ebauche|ébauche) a$", re.IGNORECASE)
_DRAFT_B = re.compile(r"^(draft|draf|ebauche|ébauche) b$", re.IGNORECASE)
_RESPONSE = re.compile(
    r"^(response|tanggapan|réponse|corriger / recommencer|correct / start over|can the ai draft help)", re.IGNORECASE
)
_CORRECTION = re.compile(r"^(correction|koreksi)$", re.IGNORECASE)
_HUMAN = re.compile(r"^human translation$", re.IGNORECASE)
_REFERENCE = re.compile(r"^\s*(\d?\s?[^\d\s:]+)\s+(\d+)(?::(\d+)(?:\s*-\s*(\d+))?)?")
_VERSE = re.compile(r"^(\d+)(?:\.0)?(?:\s*-\s*(\d+))?$")


def book_code(name: str) -> Optional[str]:
    return BOOKS.get(name.replace(" ", "").upper())


def preference(response: str) -> Optional[str]:
    """A or B for the better draft, "good" or "poor" when the team judged them equal."""
    text = response.strip().lower()
    if re.search(r"\b(draft|draf|ebauche|ébauche) a\b", text) or text == "a":
        return "A"
    if re.search(r"\b(draft|draf|ebauche|ébauche) b\b", text) or text == "b":
        return "B"
    if "good" in text or "bagus" in text or "bon" in text:
        return "good"
    if "poor" in text or "buruk" in text or "mauvais" in text:
        return "poor"
    return None


def usable(response: str) -> Optional[int]:
    """1 if the team would correct the draft (or it would help them), 0 if they would start over from the source."""
    text = response.strip().lower()
    if re.match(r"(correct|edit|perbaiki|perbiaki|corriger|yes|oui|ya)\b", text):
        return 1
    if re.match(r"(start over|mulai lagi|recommencer|no|non|tidak)\b", text):
        return 0
    return None


@dataclass(frozen=True)
class Team:
    folder: str
    language: str
    target_extract: str
    source_extract: str
    # For teams trialled apart from the rest of their trial; otherwise TRIAL_STARTS applies.
    trial_start: str = ""


class FieldTestTeams:
    def __init__(self, teams: List[Team]) -> None:
        self._teams = teams

    @classmethod
    def load(cls) -> "FieldTestTeams":
        table = pd.read_csv(Path(__file__).with_name("field_test_teams.tsv"), sep="\t", dtype=str).fillna("")
        return cls([Team(**row) for row in table.to_dict("records")])

    def all(self) -> List[Team]:
        return list(self._teams)

    def for_path(self, path: Path) -> Optional[Team]:
        """The team whose folder the workbook is in."""
        for part in reversed(path.parent.parts):
            for team in self._teams:
                if part == team.folder or part.startswith(team.folder + " ("):
                    return team
        return None


@dataclass
class _Columns:
    header_row: int
    reference: int
    source: Optional[int]
    drafts: List[int]
    response: Optional[int]
    correction: Optional[int]
    human: Optional[int]

    @property
    def kind(self) -> Optional[str]:
        if len(self.drafts) == 2 and self.response is not None:
            return "comparison"
        if len(self.drafts) == 1 and self.human is not None and self.response is not None:
            return "human_ai"
        if len(self.drafts) == 1 and self.correction is not None:
            return "correction"
        if len(self.drafts) == 1 and self.response is not None:
            return "usability"
        return None


def _find_columns(rows: List[List[str]]) -> Optional[_Columns]:
    for i, row in enumerate(rows[:25]):
        drafts_ab = [j for j, c in enumerate(row) if _DRAFT_A.match(c) or _DRAFT_B.match(c)]
        drafts = drafts_ab or [j for j, c in enumerate(row) if _DRAFT.match(c)]
        if not drafts:
            continue
        response = next((j for j, c in enumerate(row) if _RESPONSE.match(c)), None)
        correction = next((j for j, c in enumerate(row) if _CORRECTION.match(c)), None)
        human = next((j for j, c in enumerate(row) if _HUMAN.match(c)), None)
        filled = [j for j, c in enumerate(row) if c]
        reference = filled[0]
        # The source column sits between the reference and the drafts (named after the source text, e.g. "TBI").
        between = [j for j in filled if reference < j < min(drafts) and j != human]
        columns = _Columns(i, reference, between[0] if between else None, drafts[:2], response, correction, human)
        if columns.kind:
            return columns
    return None


class _VerseReader:
    """Verse references from a column holding either full references or verse numbers under a passage heading
    ("Marc 7:1-9", "MAT 17 and 20"), moving to the heading's next chapter when the verse numbers start over."""

    def __init__(self, heading: str) -> None:
        match = _REFERENCE.match(heading)
        self.book = book_code(match.group(1)) if match else None
        self.chapters = (
            [int(c) for c in re.findall(r"(?<![:\d-])(\d+)(?![:\d])", heading[match.end(1) :])] if match else []
        )
        if match and match.group(2) and not self.chapters:
            self.chapters = [int(match.group(2))]
        self._chapter_index = 0
        self._last_verse = 0

    def read(self, cell: str) -> Optional[str]:
        cell = cell.strip()
        match = _REFERENCE.match(cell)
        if match and match.group(3):
            book = book_code(match.group(1))
            if book is None:
                return None
            self.book, self.chapters, self._chapter_index = book, [int(match.group(2))], 0
            self._last_verse = int(match.group(3))
            return _vref(book, int(match.group(2)), int(match.group(3)), match.group(4))
        match = _VERSE.match(cell)
        if not match or self.book is None or not self.chapters:
            return None
        verse = int(match.group(1))
        if verse < self._last_verse and self._chapter_index + 1 < len(self.chapters):
            self._chapter_index += 1
        self._last_verse = verse
        return _vref(self.book, self.chapters[self._chapter_index], verse, match.group(2))


def _vref(book: str, chapter: int, verse: int, last: Optional[str]) -> str:
    return f"{book} {chapter}:{verse}" + (f"-{last}" if last and int(last) != verse else "")


def read_workbook(path: Path) -> Iterator[Tuple[str, str, pd.DataFrame]]:
    """(sheet name, kind, rows) for every judgment sheet in the workbook."""
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = [["" if c is None else str(c).strip() for c in row] for row in sheet.iter_rows(values_only=True)]
        columns = _find_columns(rows)
        if columns is None:
            continue
        reader = _VerseReader(rows[columns.header_row][columns.reference])

        def cell(row: List[str], j: Optional[int]) -> str:
            return row[j] if j is not None and j < len(row) else ""

        records = []
        for row in rows[columns.header_row + 1 :]:
            drafts = [cell(row, j) for j in columns.drafts]
            # Some sheets hold several passages, each under a repeated header row naming its passage.
            if _DRAFT.match(drafts[0]):
                reader = _VerseReader(cell(row, columns.reference))
                continue
            vref = reader.read(cell(row, columns.reference))
            if vref is None or not all(drafts):
                continue
            # Unanswered rows (a second group's copy that was never filled in) carry no judgment.
            if columns.kind != "correction" and not cell(row, columns.response):
                continue
            records.append(
                {
                    "vref": vref,
                    "source": cell(row, columns.source),
                    "draft_a": drafts[0],
                    "draft_b": drafts[1] if len(drafts) > 1 else "",
                    "response": cell(row, columns.response),
                    "correction": cell(row, columns.correction),
                    "human_translation": cell(row, columns.human),
                }
            )
        if records:
            yield sheet.title, columns.kind, pd.DataFrame(records)


def collect(root: Path, teams: FieldTestTeams) -> pd.DataFrame:
    """Every judged verse in the exported workbooks, one row each. Sheets copied into several workbooks (e.g. the
    human-vs-AI comparison) are kept once."""
    frames, seen = [], set()
    for path in sorted(root.rglob("*.xlsx")):
        relative = path.relative_to(root).as_posix()
        team = teams.for_path(path)
        if team is None or "Compile Feedback" in relative or "Internal Use" in relative:
            continue
        for sheet, kind, records in read_workbook(path):
            digest = hashlib.sha1(records.to_csv(index=False).encode("utf-8")).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            labels = records["response"].map(preference if kind == "comparison" else usable)
            frames.append(
                records.assign(
                    trial=relative.split("/")[0],
                    team=team.folder,
                    language=team.language,
                    target_extract=team.target_extract,
                    source_extract=team.source_extract,
                    trial_start=team.trial_start,
                    workbook=relative,
                    sheet=sheet,
                    kind=kind,
                    label=labels,
                )
            )
    table = pd.concat(frames, ignore_index=True)
    table.insert(0, "comparison", table["workbook"] + "#" + table["sheet"])
    return table


# Rows of the compiled verse-analysis sheets, by the label in their second column, and the column each becomes. For
# the yes/no and agree/disagree questions, 1 is the answer that speaks well of the draft.
ANALYSIS_QUESTIONS = {
    "hard to read aloud": ("not_hard_to_read", "N"),
    "confusing": ("not_confusing", "N"),
    "same basic meaning": ("same_meaning", "A"),
    "clear": ("clear", "A"),
    "sounds right": ("sounds_right", "A"),
}
ERROR_TYPES = [
    "wrong name",
    "wrong kinship term",
    "wrong word",
    "not a word",
    "repeated word",
    "added word",
    "missing word",
    "word order",
    "spelling",
    "wrong number",
    "wrong punctuation",
    "other",
]
ANALYSIS_LABELS = [column for column, _ in ANALYSIS_QUESTIONS.values()] + ["score", "edit_proportion"]
ANALYSIS_LABELS += ["error_" + error.replace(" ", "_") for error in ERROR_TYPES]
# Team names in the compiled sheets that differ from the folder names.
_TEAM_ALIASES = {"team3": "Team 3", "team15": "Team 15", "team20": "Team 20"}


def _rows(path: Path, sheet_filter=None) -> Iterator[Tuple[str, List[List[str]]]]:
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        if sheet_filter is None or sheet_filter(sheet.title):
            yield sheet.title, [
                ["" if c is None else str(c).strip() for c in r] for r in sheet.iter_rows(values_only=True)
            ]


def read_verse_analyses(path: Path) -> pd.DataFrame:
    """The compiled per-verse analyses: a team name, passage headings across one row and verse numbers below them,
    then a row per question, error type and score."""
    records = []
    for _, rows in _rows(path):
        if len(rows) < 8 or not any("read aloud" in (r[1] if len(r) > 1 else "") for r in rows[3:8]):
            continue
        team = next((c for c in rows[0] if c), "")
        verses: Dict[int, str] = {}
        reader = None
        for j in range(2, len(rows[2])):
            if j < len(rows[1]) and rows[1][j]:
                reader = _VerseReader(rows[1][j])
            vref = reader.read(rows[2][j]) if reader and rows[2][j] else None
            if vref:
                verses[j] = vref
        labels: Dict[int, Dict[str, float]] = {j: {} for j in verses}
        for row in rows[3:]:
            name = (row[1] if len(row) > 1 else "").lower().rstrip("?").strip()
            heading = row[0].lower() if row else ""
            if heading.startswith("correlation"):
                break
            error = re.sub(r"^\d+\.?\s*", "", name)
            for j in verses:
                value = row[j] if j < len(row) else ""
                question = next((q for q in ANALYSIS_QUESTIONS if name.startswith(q)), None)
                if question:
                    column, good = ANALYSIS_QUESTIONS[question]
                    if value.upper() in ("Y", "N", "A", "D"):
                        labels[j][column] = float(value.upper() == good)
                elif error in ERROR_TYPES:
                    labels[j]["error_" + error.replace(" ", "_")] = float(value) if value else 0.0
                elif name.startswith("score out of 100") and value:
                    labels[j]["score"] = float(value)
                elif heading.startswith("edit proportion") and value:
                    labels[j]["edit_proportion"] = float(value)
        for j, vref in verses.items():
            if "score" in labels[j] or "same_meaning" in labels[j]:
                records.append({"team_name": team, "vref": vref, **labels[j]})
    return pd.DataFrame(records)


def read_passage_texts(path: Path) -> pd.DataFrame:
    """Source and draft text of each verse in the passage tables of a verse selection or correction sheet: a heading
    row ("MAT 9:26-34", source name, "... Draft") followed by verse numbers."""
    records = []
    for title, rows in _rows(path, lambda t: "nalys" in t or "orrect" in t or "Edit" in t):
        reader, draft_column = None, None
        for row in rows:
            first = row[0] if row else ""
            drafts = [j for j, c in enumerate(row) if j > 0 and "draft" in c.lower() or c.lower().startswith("draf")]
            if drafts and _REFERENCE.match(first):
                reader, draft_column = _VerseReader(first), drafts[0]
                continue
            vref = reader.read(first) if reader and first else None
            if vref and draft_column < len(row) and row[draft_column]:
                records.append({"vref": vref, "source": row[1] if draft_column > 1 else "", "draft": row[draft_column]})
    return pd.DataFrame(records, columns=["vref", "source", "draft"])


def collect_analyses(root: Path, teams: FieldTestTeams, corrections: pd.DataFrame) -> pd.DataFrame:
    """The verse analyses with each verse's draft and the team's correction of it, where known."""
    by_name = {t.folder.lower(): t for t in teams.all()}
    analyses = []
    for path in sorted(root.rglob("Compile Feedback*.xlsx")):
        analyses.append(read_verse_analyses(path).assign(trial=path.relative_to(root).parts[0]))
    analyses = pd.concat(analyses, ignore_index=True)
    # One regional workbook repeats the main workbook's sheets for its five teams.
    analyses = analyses.drop_duplicates(["team_name", "vref"])
    texts = []
    for path in sorted(root.rglob("*.xlsx")):
        team = teams.for_path(path)
        if team and ("Verse Selections" in path.name or "Correction" in path.name):
            texts.append(read_passage_texts(path).assign(team=team.folder))
    texts = pd.concat(texts, ignore_index=True).drop_duplicates(["team", "vref"])
    records = []
    for row in analyses.to_dict("records"):
        key = row.pop("team_name").lower()
        team = by_name.get(key) or by_name.get(_TEAM_ALIASES.get(key, "").lower())
        if team is None:
            team = next((t for t in teams.all() if t.folder.lower().startswith(key)), None)
        if team is None:
            LOGGER.warning(f"No team for verse analyses of {key}")
            continue
        records.append({**row, "team": team.folder, "team_info": team})
    table = pd.DataFrame(records).merge(texts, on=["team", "vref"], how="left")
    table = table.merge(
        corrections[["team", "vref", "draft_a", "correction"]].drop_duplicates(["team", "vref"]),
        on=["team", "vref"],
        how="left",
    )
    table["draft"] = table["draft"].fillna(table["draft_a"])
    missing = table["draft"].isna().sum()
    if missing:
        LOGGER.warning(f"No draft text for {missing} analysed verses")
    table = table[table["draft"].notna()]
    info = table.pop("team_info")
    return (
        table.drop(columns="draft_a")
        .assign(
            comparison="analysis#" + table["team"],
            language=[t.language for t in info],
            target_extract=[t.target_extract for t in info],
            source_extract=[t.source_extract for t in info],
            trial_start=[t.trial_start for t in info],
            workbook="Compile Feedback",
            sheet="",
            kind="analysis",
            draft_a=table["draft"],
            draft_b="",
            response="",
            human_translation="",
            label=table["score"],
            source=table["source"].fillna(""),
            correction=table["correction"].fillna(""),
        )
        .drop(columns="draft")
    )


# When each trial's drafts were made: just after the extracts its models were trained on (xyz-ABC_2025_09_02, the
# second-trial projects' 2025_10_21-22 extracts). A reference verse already in the project before then can't have been
# written from the drafts; one that appeared later may have been, which biases it toward the draft the team imported.
TRIAL_STARTS = {"nt_field_test": "2025-09-03", "trial_2": "2025-10-23"}


def dated_extracts(scripture_dirs: Iterable[Path], project: str) -> List[Tuple[str, Path]]:
    """Every extract of a project, oldest first, with its date ("" for an undated one)."""
    extracts = []
    for directory in scripture_dirs:
        for path in directory.glob(f"{project}_20*.txt"):
            match = re.search(r"_(\d{4})_(\d\d)_(\d\d)$", path.stem)
            if match:
                extracts.append(("-".join(match.groups()), path))
        if (directory / f"{project}.txt").is_file():
            extracts.append(("", directory / f"{project}.txt"))
    return sorted(extracts, key=lambda extract: extract[0] or "9999")


def attach_references(table: pd.DataFrame, scripture_dirs: Iterable[Path], vref_path: Path) -> pd.DataFrame:
    """Each verse's text in the latest extract of the team's project ("" where the project lacks it), the date of
    the first extract that has it, and whether that predates the trial."""
    index = {v: i for i, v in enumerate(vref_path.read_text(encoding="utf-8").splitlines())}
    projects: Dict[str, List[Tuple[str, str, List[str]]]] = {}
    references, extracts, first_seen = [], [], []
    for row in table.itertuples(index=False):
        if row.target_extract not in projects:
            projects[row.target_extract] = [
                (date, path.stem, path.read_text(encoding="utf-8-sig").splitlines())
                for date, path in dated_extracts(scripture_dirs, row.target_extract)
            ]

        def text(lines: List[str]) -> str:
            verses = [
                lines[index[v]].strip() if v in index and index[v] < len(lines) else "" for v in expand_vref(row.vref)
            ]
            return " ".join(verses) if all(v and v != "<range>" for v in verses) else ""

        versions = projects[row.target_extract]
        latest = text(versions[-1][2]) if versions else ""
        references.append(latest)
        extracts.append(versions[-1][1] if versions else "")
        first_seen.append(next((date for date, _, lines in versions if text(lines)), "") if latest else "")
    table = table.assign(reference=references, reference_extract=extracts, reference_first_seen=first_seen)
    starts = table["trial"].map(TRIAL_STARTS)
    if "trial_start" in table:
        starts = table["trial_start"].where(table["trial_start"] != "", starts)
    return table.assign(
        independent_reference=(table["reference_first_seen"] != "") & (table["reference_first_seen"] < starts)
    )


def main() -> None:
    from ..features.corpora import BibleCorpus

    parser = argparse.ArgumentParser(description="Collect the field-test judgments into one table")
    parser.add_argument("--root", type=Path, default=Path("data/external/field_tests"))
    parser.add_argument("--output", type=Path, default=Path("data/field_tests/judgments.tsv"))
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    corpus = BibleCorpus()
    teams = FieldTestTeams.load()
    judgments = collect(args.root, teams)
    analyses = collect_analyses(args.root, teams, judgments[judgments["kind"] == "correction"])
    table = pd.concat([judgments, analyses], ignore_index=True)
    table = attach_references(table, corpus.scripture_dirs, corpus.vref_path)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(args.output, sep="\t", index=False)
    summary = table.assign(labelled=table["label"].notna(), has_reference=table["reference"] != "")
    columns = ["labelled", "has_reference", "independent_reference"]
    LOGGER.info(
        "Judged verses per team and kind (labelled / with a reference / with one from before the trial):\n"
        + summary.groupby(["team", "kind"])[columns].sum().astype(int).to_string()
    )
